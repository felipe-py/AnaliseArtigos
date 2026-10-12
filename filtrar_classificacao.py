#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para filtrar e salvar artigos/registros classificados com:
- Média-Alta (★☆☆)
- Alta (★★☆)
- Excelente (★★★)

Compatível com a base consolidada de mineração e sensoriamento remoto.
Gera tanto um arquivo Excel (.xlsx) formatado quanto uma cópia em CSV (.csv).
"""

import sys
import os
import re
import unicodedata
import argparse
from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Configura encoding de saída para evitar erro em consoles Windows (cp1252)
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass


def normalizar_texto(texto: str) -> str:
    """Remove acentos e converte para minúsculas para busca robusta."""
    if not isinstance(texto, str):
        return ""
    return unicodedata.normalize("NFKD", texto).encode("ASCII", "ignore").decode("utf-8").lower()


def eh_classificacao_alvo(valor) -> bool:
    """
    Verifica se a classificação corresponde a Média-Alta, Alta ou Excelente.
    Suporta variações com ícones (★), hífens, espaços e acentuações.
    Exclui categorias como Baixa, Irrelevante ou Média (isolada).
    """
    if valor is None or pd.isna(valor):
        return False
    
    norm = normalizar_texto(str(valor))
    
    # Excelente
    if "excelente" in norm:
        return True
    
    # Média-Alta ou Media Alta
    if "media-alta" in norm or "media alta" in norm or "media_alta" in norm:
        return True
    
    # Alta (garantindo que não seja Baixa, Média simples ou Irrelevante)
    # Exemplo: '★★☆ Alta' -> norm tem 'alta'
    if re.search(r"\balta\b", norm):
        return True
    
    return False


def encontrar_coluna_classificacao(df: pd.DataFrame, coluna_especificada: str = None) -> str:
    """Localiza a coluna de classificação no DataFrame."""
    if coluna_especificada and coluna_especificada in df.columns:
        return coluna_especificada

    for col in df.columns:
        norm = normalizar_texto(col)
        if "classific" in norm:
            return col
            
    raise ValueError(
        f"Coluna de classificação não encontrada. Colunas disponíveis: {list(df.columns)}"
    )


def formatar_planilha_excel(caminho_xlsx: Path):
    """
    Aplica formatação visual profissional na planilha gerada:
    - Cabeçalho com fundo azul escuro e texto branco em negrito
    - Congelamento da primeira linha (freeze panes)
    - Auto-ajuste de largura de colunas
    - Alinhamentos apropriados
    """
    try:
        wb = openpyxl.load_workbook(caminho_xlsx)
        
        # Estilos
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Segoe UI", size=10)
        border_thin = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )
        
        for ws in wb.worksheets:
            ws.views.sheetView[0].showGridLines = True
            
            # Se for a aba de dados principais
            if ws.title in ["Artigos Filtrados", "Ranking Consolidado"]:
                ws.freeze_panes = "A2"
                
                # Formata cabeçalho
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                ws.row_dimensions[1].height = 28
                
                # Formata células de dados
                for row_idx in range(2, ws.max_row + 1):
                    ws.row_dimensions[row_idx].height = 20
                    for col_idx in range(1, ws.max_column + 1):
                        cell = ws.cell(row=row_idx, column=col_idx)
                        cell.font = regular_font
                        cell.border = border_thin
                        cell.alignment = Alignment(vertical="center")
                        
                        # Formata números com 1 casa decimal se score
                        col_name = str(ws.cell(1, col_idx).value).lower()
                        if "score" in col_name and isinstance(cell.value, (int, float)):
                            cell.number_format = "0.0"
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                        elif col_name in ["rank", "ano"] and isinstance(cell.value, (int, float)):
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                
                # Ajuste de largura das colunas
                for col in ws.columns:
                    col_letter = get_column_letter(col[0].column)
                    header_val = str(col[0].value or "")
                    
                    max_len = max(len(str(cell.value or "")) for cell in col[:50])
                    # Largura estimada
                    largura = max(max_len + 3, len(header_val) + 4, 10)
                    
                    # Limita colunas muito longas (títulos/termos) para facilitar visualização
                    if largura > 60:
                        largura = 60
                    ws.column_dimensions[col_letter].width = largura

        wb.save(caminho_xlsx)
    except Exception as e:
        print(f"[Aviso] Não foi possível aplicar formatação estética avançada: {e}")


def filtrar_e_salvar(
    arquivo_entrada: str,
    arquivo_saida_xlsx: str = "artigos_media_alta_alta_excelente.xlsx",
    arquivo_saida_csv: str = "artigos_media_alta_alta_excelente.csv",
    nome_aba: str = None,
    coluna_classificacao: str = None,
    incluir_abas_apoio: bool = True,
):
    path_entrada = Path(arquivo_entrada)
    if not path_entrada.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {path_entrada}")

    print("=" * 70)
    print(" FILTRO DE ARTIGOS: MÉDIA-ALTA, ALTA E EXCELENTE")
    print("=" * 70)
    print(f"-> Arquivo de entrada: {path_entrada.resolve()}")

    # Lê o arquivo Excel
    xl = pd.ExcelFile(path_entrada)
    todas_abas = xl.sheet_names
    
    # Determina qual aba usar
    aba_dados = nome_aba
    if not aba_dados:
        if "Ranking Consolidado" in todas_abas:
            aba_dados = "Ranking Consolidado"
        else:
            aba_dados = todas_abas[0]

    print(f"-> Aba analisada: '{aba_dados}'")
    df_original = xl.parse(aba_dados)
    total_linhas_originais = len(df_original)
    print(f"-> Total de registros na base original: {total_linhas_originais:,}")

    # Encontra coluna de classificação
    col_class = encontrar_coluna_classificacao(df_original, coluna_classificacao)
    print(f"-> Coluna de classificação identificada: '{col_class}'")

    # Aplica o filtro
    mascara = df_original[col_class].apply(eh_classificacao_alvo)
    df_filtrado = df_original[mascara].copy()
    total_filtrado = len(df_filtrado)

    print("\n" + "-" * 70)
    print(f" RESULTADO DO FILTRO: {total_filtrado} registros encontrados ({total_filtrado/total_linhas_originais*100:.2f}%)")
    print("-" * 70)
    
    distribuicao = df_filtrado[col_class].value_counts()
    for cat, qtd in distribuicao.items():
        print(f"   • {cat:<25}: {qtd:>4} artigos ({qtd/total_filtrado*100:5.1f}%)")

    # Identifica coluna de score para estatísticas extras (se houver)
    col_score = None
    for c in df_filtrado.columns:
        if "score" in c.lower() and ("100" in c or "total" in c or c.lower().strip() == "score"):
            col_score = c
            break
            
    if col_score:
        print(f"\n Estatísticas de Score ({col_score}):")
        print(f"   • Mínimo : {df_filtrado[col_score].min():.1f}")
        print(f"   • Média  : {df_filtrado[col_score].mean():.1f}")
        print(f"   • Máximo : {df_filtrado[col_score].max():.1f}")

    # Prepara resumo para salvar no Excel
    resumo_dados = {
        "Métrica / Categoria": [
            "Total original de artigos",
            "Total filtrado (Média-Alta, Alta, Excelente)",
            "Percentual retido",
            "---",
        ] + [f"Categoria: {cat}" for cat in distribuicao.index],
        "Quantidade": [
            total_linhas_originais,
            total_filtrado,
            f"{total_filtrado/total_linhas_originais*100:.2f}%",
            "---",
        ] + list(distribuicao.values),
    }
    df_resumo = pd.DataFrame(resumo_dados)

    # 1. Salva em XLSX
    path_saida_xlsx = Path(arquivo_saida_xlsx)
    print("\n" + "-" * 70)
    print(f"-> Gerando arquivo Excel: {path_saida_xlsx.resolve()}")
    
    with pd.ExcelWriter(path_saida_xlsx, engine="openpyxl") as writer:
        # Aba principal com os artigos filtrados
        df_filtrado.to_excel(writer, sheet_name="Artigos Filtrados", index=False)
        # Aba com o resumo estatístico do filtro
        df_resumo.to_excel(writer, sheet_name="Resumo do Filtro", index=False)
        
        # Se solicitado, copia as abas de critérios e glossários da base original
        if incluir_abas_apoio:
            for aba in todas_abas:
                if aba != aba_dados:
                    try:
                        df_apoio = xl.parse(aba)
                        # Limita nome da aba a 31 caracteres (limite do Excel)
                        nome_aba_apoio = aba[:31]
                        df_apoio.to_excel(writer, sheet_name=nome_aba_apoio, index=False)
                    except Exception as err:
                        print(f"[Aviso] Não foi possível copiar aba '{aba}': {err}")

    # Aplica formatação estética
    formatar_planilha_excel(path_saida_xlsx)
    print("   ✓ Arquivo Excel gerado com formatação e abas preservadas.")

    # 2. Salva em CSV (se solicitado)
    if arquivo_saida_csv:
        path_saida_csv = Path(arquivo_saida_csv)
        print(f"-> Gerando arquivo CSV  : {path_saida_csv.resolve()}")
        # utf-8-sig garante abertura perfeita de acentos e símbolos no Excel e ferramentas externas
        df_filtrado.to_csv(path_saida_csv, index=False, encoding="utf-8-sig", sep=";")
        print("   ✓ Arquivo CSV gerado com codificação UTF-8 com BOM (separador ';').")

    print("=" * 70)
    print(" CONCLUÍDO COM SUCESSO!")
    print("=" * 70)
    return df_filtrado


def main():
    parser = argparse.ArgumentParser(
        description="Filtra artigos com classificação Média-Alta, Alta ou Excelente de planilha Excel."
    )
    parser.add_argument(
        "--input", "-i",
        default="base_consolidada_ranking_mineracao_DL_SR_perguntas.xlsx",
        help="Caminho do arquivo Excel de entrada (padrão: base_consolidada_ranking_mineracao_DL_SR_perguntas.xlsx)",
    )
    parser.add_argument(
        "--output", "-o",
        default="artigos_media_alta_alta_excelente.xlsx",
        help="Caminho do arquivo Excel de saída (padrão: artigos_media_alta_alta_excelente.xlsx)",
    )
    parser.add_argument(
        "--csv",
        default="artigos_media_alta_alta_excelente.csv",
        help="Caminho do arquivo CSV de saída (padrão: artigos_media_alta_alta_excelente.csv). Passe '' para desabilitar.",
    )
    parser.add_argument(
        "--sheet", "-s",
        default=None,
        help="Nome da aba a ser filtrada (padrão: detecta automaticamente)",
    )
    parser.add_argument(
        "--column", "-c",
        default=None,
        help="Nome da coluna de classificação (padrão: detecta automaticamente)",
    )
    parser.add_argument(
        "--no-ref-sheets",
        action="store_true",
        help="Não copiar abas de apoio/critérios no Excel de saída",
    )

    args = parser.parse_args()

    # Se o arquivo padrão não existir no diretório atual, procura qualquer .xlsx no diretório
    arquivo_input = args.input
    if not os.path.exists(arquivo_input):
        xlsx_candidatos = list(Path(".").glob("*.xlsx"))
        if xlsx_candidatos:
            arquivo_input = str(xlsx_candidatos[0])
            print(f"[Info] Usando arquivo encontrado: {arquivo_input}")

    csv_output = args.csv if args.csv != "" else None

    filtrar_e_salvar(
        arquivo_entrada=arquivo_input,
        arquivo_saida_xlsx=args.output,
        arquivo_saida_csv=csv_output,
        nome_aba=args.sheet,
        coluna_classificacao=args.column,
        incluir_abas_apoio=not args.no_ref_sheets,
    )


if __name__ == "__main__":
    main()
